"""Core logic: read a file, decide which project it belongs to, move it, and log the move.

Nothing in here is Windows specific, so it can be tested on any computer.
Every function has exactly one return statement.
"""
import json
import os
import re
import shutil
import time
import urllib.request
from datetime import datetime
from pathlib import Path

# Files that should never be moved automatically
SKIP_SUFFIXES = {".lnk", ".url", ".ini", ".tmp", ".crdownload", ".part", ".partial", ".download", ".opdownload"}
SKIP_NAME_STARTS = ("~$", ".", "New Text Document", "New Microsoft", "New Rich Text", "New Bitmap")

# Files whose text we can read directly
TEXT_SUFFIXES = {".txt", ".md", ".csv", ".json", ".py", ".js", ".ts", ".cs", ".cpp", ".c", ".h", ".hpp",
                 ".php", ".html", ".sql", ".r", ".ino", ".java", ".xml", ".log", ".ipynb"}
MAX_TEXT = 3000
HIDDEN_FLAG = 2
LINK_FLAG = 0x400  # symbolic links and junctions such as "My Music" inside Documents

# Signs that a folder is a programming project. Those folders are left exactly as they are.
CODE_MARKERS = {".git", ".svn", ".hg", ".vscode", ".idea", "node_modules", "venv", ".venv", "__pycache__",
                ".gitignore", "package.json", "requirements.txt", "pyproject.toml", "setup.py", "pipfile",
                "cargo.toml", "go.mod", "pom.xml", "build.gradle", "cmakelists.txt", "makefile", "gemfile",
                "composer.json", "platformio.ini"}
CODE_MARKER_SUFFIXES = {".sln", ".csproj", ".vcxproj", ".xcodeproj", ".uproject", ".code-workspace"}
CODE_SUFFIXES = {".py", ".js", ".ts", ".jsx", ".tsx", ".cs", ".cpp", ".c", ".h", ".hpp", ".java", ".kt", ".go",
                 ".rs", ".rb", ".php", ".swift", ".ino", ".ipynb", ".sh", ".ps1", ".lua", ".r", ".scala", ".vue",
                 ".css", ".scss", ".sql"}
# Files a program needs next to it. A folder holding any of these belongs to a program and is left alone.
APP_SUFFIXES = {".dll", ".so", ".dylib", ".sys", ".pak", ".asar", ".sqlite", ".ldb", ".pst", ".ost"}
# Folders that programs keep in your Documents for themselves
LEAVE_ALONE = {"zoom", "custom office templates", "my games", "windowspowershell", "powershell", "outlook files",
               "onenote notebooks", "visual studio 2019", "visual studio 2022", "adobe", "fax", "scanned documents"}

# Files that can point at other files by name: a web page and its pictures, a playlist and its songs
REFERENCE_SUFFIXES = {".html", ".htm", ".xhtml", ".css", ".svg", ".md", ".tex", ".m3u", ".m3u8", ".pls", ".cue",
                      ".xspf"}
MAX_REFERENCE_TEXT = 200_000
# Names too common to mean two files belong together just because they share them
GENERIC_NAMES = {"image", "download", "untitled", "document", "file", "new", "copy", "scan", "export"}


# ---------------------------------------------------------------- places we keep our own data

def data_folder():
    """Folder for the move log and the ignore list (outside the Projects folder so it stays clean)."""
    base = os.environ.get("LOCALAPPDATA") or str(Path.home() / ".local" / "share")
    folder = Path(base) / "FileOrganizer"
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def ensure_config(config_path):
    """On first run, start your own config.json from config.example.json."""
    example = config_path.with_name("config.example.json")
    created = not config_path.exists() and example.exists()
    if created:
        shutil.copyfile(example, config_path)
    return created


def load_config(config_path):
    with open(config_path, encoding="utf-8") as handle:
        config = json.load(handle)
    return config


# ---------------------------------------------------------------- which files to leave alone

def load_ignored():
    """Files the user moved back with Undo. We never auto sort these again."""
    ignore_file = data_folder() / "ignored.json"
    ignored = set()
    if ignore_file.exists():
        try:
            ignored = set(json.loads(ignore_file.read_text(encoding="utf-8")))
        except (ValueError, OSError):
            ignored = set()
    return ignored


def save_ignored(ignored):
    ignore_file = data_folder() / "ignored.json"
    ignore_file.write_text(json.dumps(sorted(ignored), indent=2), encoding="utf-8")
    return len(ignored)


def is_open_in_office(path):
    """Word and Excel create a '~$name' file next to a document while it is open."""
    lock_name = "~$" + path.name[2:]
    return (path.parent / lock_name).exists()


def is_hidden(path):
    attributes = getattr(path.stat(), "st_file_attributes", 0)
    return bool(attributes & HIDDEN_FLAG)


def is_sortable(path, automatic=True):
    """True when a file is safe to move. Folders, shortcuts, hidden and half downloaded files are skipped."""
    sortable = False
    try:
        sortable = (path.is_file()
                    and path.suffix.lower() not in SKIP_SUFFIXES
                    and not path.name.startswith(SKIP_NAME_STARTS)
                    and not is_hidden(path)
                    and not is_open_in_office(path))
        if sortable and automatic:
            sortable = str(path) not in load_ignored()
    except OSError:
        sortable = False
    return sortable


def is_link(path):
    attributes = getattr(path.lstat(), "st_file_attributes", 0)
    return path.is_symlink() or bool(attributes & LINK_FLAG)


def folder_kind(folder):
    """'code' for a programming project, 'app' for a program's own files, otherwise 'plain' (safe to tidy)."""
    kind = "plain"
    try:
        children = list(folder.iterdir())
        names = {child.name.lower() for child in children}
        file_suffixes = [child.suffix.lower() for child in children if child.is_file()]
        code_files = sum(1 for suffix in file_suffixes if suffix in CODE_SUFFIXES)
        if (names & CODE_MARKERS
                or CODE_MARKER_SUFFIXES.intersection(child.suffix.lower() for child in children)
                or (code_files >= 2 and code_files * 2 >= len(file_suffixes))):
            kind = "code"
        elif APP_SUFFIXES.intersection(file_suffixes):
            kind = "app"
    except OSError:
        kind = "app"  # we cannot look inside, so we leave it alone
    return kind


# ---------------------------------------------------------------- finding files that depend on each other

def shared_name(path):
    """The part of a name related files share: 'trip.mp4', 'trip.srt', 'trip.part2.rar' and 'trip_files' all give 'trip'.
    Returns None for plain folders and very common names."""
    name = None
    if path.is_dir():
        if path.name.lower().endswith("_files"):
            name = path.name.lower()[:-len("_files")]
    else:
        name = re.sub(r"(\.part\d+|\.tar|\.zip|\.7z|\.rar)+$", "", path.stem.lower())
    if name in GENERIC_NAMES or not name:
        name = None
    return name


def mentioned_names(path, neighbours):
    """Names of neighbouring files and folders that this file mentions (neighbours maps lowercase name -> path)."""
    found = set()
    if path.suffix.lower() in REFERENCE_SUFFIXES:
        try:
            with open(path, encoding="utf-8", errors="ignore") as handle:
                text = handle.read(MAX_REFERENCE_TEXT).lower().replace("%20", " ")
        except OSError:
            text = ""
        for name, neighbour in neighbours.items():
            # a folder only counts when it is used as part of a path, like "images/logo.png"
            ending = r"[/\\]" if neighbour.is_dir() else r"(?![\w-])"
            if neighbour != path and re.search(r"(?<![\w.-])" + re.escape(name) + ending, text):
                found.add(name)
    return found


def related_groups(entries):
    """Split files (and the folders they use) into groups that depend on each other and must move together.
    Groups made only of folders are dropped, those folders are not part of anything."""
    neighbours = {entry.name.lower(): entry for entry in entries}
    leader = {entry: entry for entry in entries}

    def find(entry):
        while leader[entry] != entry:
            entry = leader[entry]
        return entry

    def join(first, second):
        leader[find(first)] = find(second)
        return None

    by_shared = {}
    for entry in entries:
        key = shared_name(entry)
        if key:
            by_shared.setdefault(key, []).append(entry)
    for members in by_shared.values():
        for other in members[1:]:
            join(members[0], other)
    for entry in entries:
        if entry.is_file():
            for name in mentioned_names(entry, neighbours):
                join(entry, neighbours[name])
    groups = {}
    for entry in entries:
        groups.setdefault(find(entry), []).append(entry)
    return [sorted(group) for group in groups.values() if any(member.is_file() for member in group)]


def find_groups(folder, config, protected, depth=0):
    """Every group of related files to sort in this folder and, when sort_inside_folders is on, in its everyday
    subfolders too. Programming projects, program folders, hidden folders and protected folders are skipped."""
    inside = config.get("sort_inside_folders", True)
    max_depth = config.get("max_folder_depth", 3) if inside else 0
    leave_alone = LEAVE_ALONE | {name.lower() for name in config.get("leave_folders_alone", [])}
    ignored = load_ignored()
    files, folders, kinds = [], [], {}
    try:
        children = sorted(folder.iterdir())
    except OSError:
        children = []
    for child in children:
        try:
            if child.is_dir():
                usable = (inside and not is_link(child) and not is_hidden(child)
                          and child.name.lower() not in leave_alone and str(child) not in ignored
                          and child.resolve() not in protected)
                if usable:
                    kinds[child] = folder_kind(child)
                    # a saved web page's "_files" folder always travels with its page, other folders only when plain
                    if kinds[child] == "plain" or child.name.lower().endswith("_files"):
                        folders.append(child)
            elif is_sortable(child, automatic=False):
                files.append(child)
        except OSError:
            pass
    groups = related_groups(files + folders)
    claimed = {member for group in groups for member in group}
    for subfolder in folders:
        if subfolder not in claimed and kinds[subfolder] == "plain" and depth < max_depth:
            groups += find_groups(subfolder, config, protected, depth + 1)
    return groups


def group_for(path, config, protected):
    """The file together with the neighbours it depends on (for the right click menu)."""
    deepest = config.get("max_folder_depth", 3)
    groups = [group for group in find_groups(path.parent, config, protected, depth=deepest) if path in group]
    return groups[0] if groups else [path]


def group_is_sortable(group):
    """Every file in the group can be moved automatically (none is open or was put back with Undo)."""
    ignored = load_ignored()
    return all(str(member) not in ignored and (member.is_dir() or is_sortable(member, automatic=True))
               for member in group)


def last_changed(group):
    """When anything in the group (including inside its folders) was last changed."""
    times = [0.0]
    for member in group:
        times.append(member.stat().st_mtime)
        if member.is_dir():
            times += [inner.stat().st_mtime for inner in member.rglob("*")]
    return max(times)


def main_file(group):
    """The file the others hang off: a page or playlist that mentions them, otherwise the biggest file."""
    files = [member for member in group if member.is_file()]
    referencing = REFERENCE_SUFFIXES - {".css", ".svg"}
    return max(files, key=lambda path: (path.suffix.lower() in referencing, path.stat().st_size))


def remove_empty_folders(folder, stop_at):
    """After files were moved out of a subfolder, remove it (and empty folders above it). Never removes stop_at."""
    removed = 0
    current = folder
    try:
        while current != stop_at and stop_at in current.parents and not any(current.iterdir()):
            current.rmdir()
            removed += 1
            current = current.parent
    except OSError:
        pass
    return removed


# ---------------------------------------------------------------- reading what is inside a file

def extract_text(path):
    """Return the first few thousand characters of text inside the file, or an empty string."""
    suffix = path.suffix.lower()
    text = ""
    try:
        if suffix in TEXT_SUFFIXES:
            text = path.read_text(encoding="utf-8", errors="ignore")
        elif suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            text = "\n".join((page.extract_text() or "") for page in reader.pages[:3])
        elif suffix == ".docx":
            import docx
            document = docx.Document(str(path))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        elif suffix == ".pptx":
            from pptx import Presentation
            deck = Presentation(str(path))
            text = "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if hasattr(shape, "text"))
    except Exception:
        text = ""
    return text[:MAX_TEXT]


# ---------------------------------------------------------------- choosing a project with simple rules

def keyword_pattern(keyword):
    """Match a keyword as its own word, so 'lab' matches 'lab3' but not 'label'."""
    return re.compile(r"(?<![a-z])" + re.escape(keyword.lower()) + r"(?![a-z])")


def keyword_scores(file_name, suffix, text, projects):
    """Give every project points: 3 per keyword in the file name, 1 per keyword in the text (max 5), 3 for a matching file type."""
    name_lower = file_name.lower()
    text_lower = text.lower()
    scores = {}
    for project, info in projects.items():
        score = 0
        for keyword in info.get("keywords", []):
            pattern = keyword_pattern(keyword)
            score += 3 * len(pattern.findall(name_lower))
            score += min(len(pattern.findall(text_lower)), 5)
        if suffix.lower() in [ext.lower() for ext in info.get("extensions", [])]:
            score += 3
        if score > 0:
            scores[project] = score
    return scores


def pick_by_rules(scores, min_score):
    """Pick the top project only if it scored enough and clearly beat the runner up."""
    choice = None
    if scores:
        ranked = sorted(scores.items(), key=lambda item: item[1], reverse=True)
        best_project, best_score = ranked[0]
        runner_up = ranked[1][1] if len(ranked) > 1 else 0
        if best_score >= min_score and best_score > runner_up:
            choice = best_project
    return choice


# ---------------------------------------------------------------- choosing a project with AI

def match_project_name(answer, project_names):
    """Turn the AI's reply into one of our project names, or None."""
    cleaned = answer.strip().strip(".\"'*` ").lower()
    choice = None
    for name in project_names:
        if cleaned == name.lower():
            choice = name
    if choice is None and "none" not in cleaned.split():
        for name in sorted(project_names, key=len, reverse=True):
            if choice is None and name.lower() in cleaned:
                choice = name
    return choice


def build_prompt(file_name, text, projects):
    project_lines = "\n".join(f"{name}: {info.get('description', '')}" for name, info in projects.items())
    content = text.strip() or "(no readable text, decide from the file name)"
    prompt = (
        "You are sorting a file from someone's computer into one of their project folders.\n\n"
        f"Projects:\n{project_lines}\n\n"
        f"File name: {file_name}\n"
        f"Start of the file's content:\n\"\"\"\n{content[:2000]}\n\"\"\"\n\n"
        "Reply with only the project name exactly as written above. "
        "If no project clearly fits, reply NONE."
    )
    return prompt


def ask_ai(file_name, text, projects, ai_settings):
    """Ask a local Ollama model. Returns a project name, or None if AI is off, unavailable, or unsure."""
    choice = None
    if ai_settings.get("enabled"):
        body = json.dumps({
            "model": ai_settings.get("model", "llama3.2:3b"),
            "prompt": build_prompt(file_name, text, projects),
            "stream": False,
            "options": {"temperature": 0},
        }).encode("utf-8")
        url = ai_settings.get("url", "http://localhost:11434").rstrip("/") + "/api/generate"
        request = urllib.request.Request(url, data=body, headers={"Content-Type": "application/json"})
        try:
            with urllib.request.urlopen(request, timeout=ai_settings.get("timeout_seconds", 60)) as response:
                answer = json.loads(response.read().decode("utf-8")).get("response", "")
            choice = match_project_name(answer, list(projects))
        except Exception:
            choice = None
    return choice


def classify(path, config):
    """Rules first, AI second, To Review if neither is sure. Returns (project, how it was decided)."""
    projects = config["projects"]
    text = extract_text(path)
    scores = keyword_scores(path.stem, path.suffix, text, projects)
    project = pick_by_rules(scores, config.get("min_keyword_score", 2))
    method = "rules"
    if project is None:
        project = ask_ai(path.name, text, projects, config.get("ai", {}))
        method = "ai"
    if project is None:
        project = config.get("review_folder", "To Review")
        method = "review"
    return project, method


# ---------------------------------------------------------------- moving and logging

def unique_path(target):
    """If a file with that name already exists, add (2), (3), and so on."""
    candidate = target
    counter = 2
    while candidate.exists():
        candidate = target.with_name(f"{target.stem} ({counter}){target.suffix}")
        counter += 1
    return candidate


def destination_for(group, main, project, root):
    """Projects/<Project>/<Year>/<Month>/, based on when the main file was last changed (To Review too).
    Related files get their own folder named after the main file, so they stay together and keep working."""
    when = datetime.fromtimestamp(main.stat().st_mtime)
    folder = root / project / f"{when:%Y}" / f"{when:%m}"
    if len(group) > 1:
        folder = unique_path(folder / re.sub(r"(\.part\d+|\.tar)+$", "", main.stem, flags=re.IGNORECASE))
    return folder


def log_move(source, target, project, method, batch, in_group_folder):
    entry = {"time": datetime.now().isoformat(timespec="seconds"), "from": str(source),
             "to": str(target), "project": project, "method": method, "batch": batch,
             "group_folder": in_group_folder}
    with open(data_folder() / "moves.jsonl", "a", encoding="utf-8") as handle:
        handle.write(json.dumps(entry) + "\n")
    return entry


def sort_group(group, config, root, forced_project=None):
    """Move a group of related files into one project, together. Returns a summary of the main file's move,
    with 'related' saying how many files or folders went along with it."""
    main = main_file(group)
    if forced_project:
        project, method = forced_project, "manual"
        ignored = load_ignored()
        ignored.difference_update(str(member) for member in group)
        save_ignored(ignored)
    else:
        project, method = classify(main, config)
    folder = destination_for(group, main, project, root)
    folder.mkdir(parents=True, exist_ok=True)
    batch = str(time.time_ns())
    summary = None
    for member in group:
        target = unique_path(folder / member.name)
        shutil.move(str(member), str(target))
        entry = log_move(member, target, project, method, batch, len(group) > 1)
        if member == main:
            summary = dict(entry, related=len(group) - 1)
    return summary


# ---------------------------------------------------------------- undo

def read_log():
    log_file = data_folder() / "moves.jsonl"
    entries = []
    if log_file.exists():
        for line in log_file.read_text(encoding="utf-8").splitlines():
            if line.strip():
                entries.append(json.loads(line))
    return entries


def undo_last(count=1):
    """Put the last few moves back where they were, and stop auto sorting those files.
    Related files that were moved together count as one move."""
    entries = read_log()
    batches = []
    for entry in entries:
        key = entry.get("batch", entry["time"] + entry["to"])  # older log lines have no batch
        if batches and batches[-1][0] == key:
            batches[-1][1].append(entry)
        else:
            batches.append((key, [entry]))
    count = max(0, min(count, len(batches)))
    kept = [entry for _, batch in batches[:len(batches) - count] for entry in batch]
    to_undo = [entry for _, batch in batches[len(batches) - count:] for entry in batch]
    ignored = load_ignored()
    restored = []
    for entry in reversed(to_undo):
        source, target = Path(entry["from"]), Path(entry["to"])
        if target.exists() and not source.exists():
            source.parent.mkdir(parents=True, exist_ok=True)
            shutil.move(str(target), str(source))
            ignored.add(str(source))
            restored.append(source)
            if entry.get("group_folder") and not any(target.parent.iterdir()):
                target.parent.rmdir()
    save_ignored(ignored)
    with open(data_folder() / "moves.jsonl", "w", encoding="utf-8") as handle:
        handle.writelines(json.dumps(entry) + "\n" for entry in kept)
    return restored


# ---------------------------------------------------------------- recent work

def recent_files(root, project_names, days, per_project):
    """For each project, the files changed most recently (newest first)."""
    cutoff = time.time() - days * 86400
    result = {}
    for name in project_names:
        folder = root / name
        files = []
        if folder.exists():
            for path in folder.rglob("*"):
                try:
                    if path.is_file() and is_sortable(path, automatic=False) and path.stat().st_mtime >= cutoff:
                        files.append(path)
                except OSError:
                    pass
        files.sort(key=lambda path: path.stat().st_mtime, reverse=True)
        if files:
            result[name] = files[:per_project]
    return result
